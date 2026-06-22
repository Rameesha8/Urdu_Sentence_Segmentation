import html
from io import BytesIO
import re

import pandas as pd
import streamlit as st
import streamlit.components.v1 as components
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from segmentation_utils import reference_split, run_all_algorithms

try:
    from dotenv import load_dotenv
except Exception:  # pragma: no cover - optional dependency
    load_dotenv = None

if load_dotenv is not None:
    load_dotenv()

st.set_page_config(
    page_title="Sentence Segmentation Lab",
    page_icon="S",
    layout="wide",
    initial_sidebar_state="expanded",
)


def inject_styles() -> None:
    st.markdown(
        """
        <style>
        :root {
            --bg: #07111f;
            --panel: rgba(10, 18, 33, 0.78);
            --panel-strong: rgba(16, 27, 48, 0.92);
            --border: rgba(255, 255, 255, 0.10);
            --text: #edf2ff;
            --muted: #a8b4cc;
            --accent: #f6c667;
            --accent-2: #57d6c2;
            --danger: #ff7a7a;
            --success: #66e3a8;
        }

        .stApp {
            background:
                radial-gradient(circle at top left, rgba(87, 214, 194, 0.18), transparent 28%),
                radial-gradient(circle at top right, rgba(246, 198, 103, 0.14), transparent 24%),
                linear-gradient(135deg, #050b14 0%, #09192d 48%, #101624 100%);
            color: var(--text);
        }

        .block-container {
            padding-top: 1.8rem;
            padding-bottom: 2.5rem;
            max-width: 1280px;
        }

        .hero {
            position: relative;
            padding: 2rem 2rem 1.5rem;
            border: 1px solid var(--border);
            border-radius: 28px;
            background: linear-gradient(160deg, rgba(255,255,255,0.08), rgba(255,255,255,0.03));
            box-shadow: 0 30px 80px rgba(0, 0, 0, 0.35);
            overflow: hidden;
        }

        .hero::after {
            content: "";
            position: absolute;
            inset: -30% -10% auto auto;
            width: 220px;
            height: 220px;
            background: radial-gradient(circle, rgba(246, 198, 103, 0.34), transparent 68%);
            filter: blur(10px);
            pointer-events: none;
        }

        .eyebrow {
            display: inline-flex;
            align-items: center;
            gap: 0.5rem;
            padding: 0.35rem 0.75rem;
            border-radius: 999px;
            border: 1px solid rgba(255,255,255,0.16);
            background: rgba(255,255,255,0.06);
            color: var(--accent);
            font-size: 0.82rem;
            letter-spacing: 0.08em;
            text-transform: uppercase;
        }

        .hero h1 {
            margin: 0.7rem 0 0.35rem;
            font-size: clamp(2.2rem, 4.3vw, 4.4rem);
            line-height: 1.03;
            color: #ffffff;
        }

        .hero p {
            margin: 0;
            color: var(--muted);
            font-size: 1.02rem;
            max-width: 900px;
        }

        .metric-card, .panel {
            border: 1px solid var(--border);
            border-radius: 22px;
            background: linear-gradient(180deg, var(--panel), rgba(10, 18, 33, 0.96));
            box-shadow: 0 18px 55px rgba(0, 0, 0, 0.28);
        }

        .metric-card {
            padding: 1rem 1.1rem;
            min-height: 104px;
        }

        .metric-label {
            color: var(--muted);
            font-size: 0.82rem;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            margin-bottom: 0.45rem;
        }

        .metric-value {
            color: #fff;
            font-size: 1.9rem;
            font-weight: 700;
            line-height: 1;
        }

        .metric-subtext {
            color: var(--muted);
            margin-top: 0.45rem;
            font-size: 0.9rem;
        }

        .winner-pill {
            display: inline-flex;
            align-items: center;
            justify-content: center;
            width: fit-content;
            margin-bottom: 0.5rem;
            padding: 0.22rem 0.6rem;
            border-radius: 999px;
            background: rgba(246, 198, 103, 0.16);
            border: 1px solid rgba(246, 198, 103, 0.35);
            color: #f6c667;
            font-size: 0.72rem;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: 0.08em;
        }

        .panel {
            padding: 1.2rem 1.25rem;
        }

        .section-title {
            margin: 0 0 0.85rem;
            font-size: 1.1rem;
            color: #fff;
        }

        .result-table {
            width: 100%;
            border-collapse: separate;
            border-spacing: 0 12px;
        }

        .result-table th {
            color: var(--muted);
            text-align: left;
            padding: 0 1rem 0.55rem 1rem;
            font-size: 0.8rem;
            letter-spacing: 0.08em;
            text-transform: uppercase;
        }

        .result-table td {
            padding: 1rem;
            vertical-align: top;
            color: var(--text);
            background: rgba(255,255,255,0.04);
            border-top: 1px solid rgba(255,255,255,0.08);
            border-bottom: 1px solid rgba(255,255,255,0.08);
        }

        .result-table td:first-child {
            border-left: 1px solid rgba(255,255,255,0.08);
            border-top-left-radius: 18px;
            border-bottom-left-radius: 18px;
            font-weight: 700;
            width: 16%;
        }

        .result-table td:last-child {
            border-right: 1px solid rgba(255,255,255,0.08);
            border-top-right-radius: 18px;
            border-bottom-right-radius: 18px;
            width: 34%;
        }

        .result-cell {
            color: #f5f8ff;
            line-height: 1.65;
            word-break: break-word;
        }

        .separator {
            color: rgba(246, 198, 103, 0.95);
            font-weight: 700;
        }

        .badge {
            display: inline-flex;
            align-items: center;
            gap: 0.4rem;
            padding: 0.32rem 0.72rem;
            border-radius: 999px;
            font-size: 0.8rem;
            font-weight: 700;
            white-space: nowrap;
        }

        .badge.good {
            color: #052415;
            background: rgba(102, 227, 168, 0.95);
        }

        .badge.review {
            color: #2d1d05;
            background: rgba(246, 198, 103, 0.95);
        }

        .badge.error {
            color: #2d0909;
            background: rgba(255, 122, 122, 0.95);
        }

        .reason-text {
            margin-top: 0.7rem;
            color: #d7def0;
            line-height: 1.55;
        }

        .result-table tbody tr:hover td {
            background: rgba(255,255,255,0.06);
        }

        .footer-note {
            color: var(--muted);
            font-size: 0.88rem;
            padding-top: 0.4rem;
        }

        div[data-testid="stFileUploader"] {
            border-radius: 18px;
            border: 1px dashed rgba(255,255,255,0.20);
            padding: 0.15rem 0.2rem;
            background: rgba(255,255,255,0.03);
        }

        textarea {
            border-radius: 16px !important;
        }

        @media (max-width: 768px) {
            .hero { padding: 1.4rem; }
            .result-table th, .result-table td { padding: 0.8rem; }
        }
        </style>
        """,
        unsafe_allow_html=True,
    )


def load_sentence() -> str:
    if "sentence_text" not in st.session_state:
        st.session_state.sentence_text = ""

    if st.button("Load sample sentence", use_container_width=True):
        st.session_state.sentence_text = "\u06cc\u06c1 \u0627\u06cc\u06a9 \u0645\u062b\u0627\u0644 \u062c\u0645\u0644\u06c1 \u06c1\u06d4 \u0627\u0633 \u0645\u06cc\u06ba \u062f\u0648 \u062d\u0635\u06d2 \u06c1\u0648 \u0633\u06a9\u062a\u06d2 \u06c1\u06cc\u06ba!"

    uploaded_file = st.file_uploader(
        "Upload a text file",
        type=["txt"],
        help="Upload a text file containing a sentence or a short paragraph.",
    )

    typed_text = st.text_area(
        "Or paste a sentence",
        key="sentence_text",
        height=160,
        placeholder="Paste Urdu or English text here...",
    )

    if uploaded_file is not None:
        return uploaded_file.read().decode("utf-8", errors="ignore").strip()
    return typed_text.strip()


def split_input_sentences(text: str) -> list[str]:
    lines = [line.strip() for line in re.split(r"[\r\n]+", text or "") if line.strip()]
    if len(lines) > 1:
        candidates = []
        for line in lines:
            candidates.extend([segment.strip() for segment in re.split(r"[.\u06d4!?]+", line) if segment.strip()])
        return candidates
    return [segment.strip() for segment in re.split(r"[.\u06d4!?]+", text or "") if segment.strip()]


def analyze_input_text(input_text: str) -> tuple[list[str], list[dict], pd.DataFrame, pd.DataFrame, str, int, int]:
    sentences = split_input_sentences(input_text)
    sentence_reports: list[dict] = []
    detailed_rows: list[dict] = []
    algorithm_names = ["Rule-Based", "Regex", "UrduHack", "Stanza", "Dataset-Trained Punkt", "Groq LLM", "Hybrid"]

    for sentence_index, sentence in enumerate(sentences, start=1):
        results = run_all_algorithms(sentence)
        best_item = max(results, key=lambda item: item["score"]) if results else None
        winner_algorithm = best_item["algorithm"] if best_item else "N/A"
        sentence_reports.append(
            {
                "sentence_index": sentence_index,
                "sentence": sentence,
                "results": results,
                "best_algorithm": winner_algorithm,
                "best_score": best_item["score"] if best_item else 0,
                "worked_well": sum(1 for item in results if item["verdict"] == "Worked well"),
                "needs_review": sum(1 for item in results if item["verdict"] == "Needs review"),
                "unavailable": sum(1 for item in results if item["verdict"] == "Could not run"),
            }
        )

        for item in results:
            detailed_rows.append(
                {
                    "sentence_index": sentence_index,
                    "sentence": sentence,
                    "algorithm": item["algorithm"],
                    "result": item["result"],
                    "verdict": item["verdict"],
                    "reason": item["reason"],
                    "score": item["score"],
                    "similarity": item["similarity"],
                    "is_winner": item["algorithm"] == winner_algorithm,
                }
            )

    detail_frame = pd.DataFrame(detailed_rows)
    summary_rows = []
    for algorithm_name in algorithm_names:
        algorithm_frame = detail_frame[detail_frame["algorithm"] == algorithm_name] if not detail_frame.empty else pd.DataFrame()
        if algorithm_frame.empty:
            summary_rows.append(
                {
                    "algorithm": algorithm_name,
                    "avg_score": 0,
                    "worked_well": 0,
                    "partially_worked": 0,
                    "needs_review": 0,
                    "could_not_run": 0,
                }
            )
            continue

        summary_rows.append(
            {
                "algorithm": algorithm_name,
                "avg_score": round(float(algorithm_frame["score"].mean()), 1),
                "worked_well": int((algorithm_frame["verdict"] == "Worked well").sum()),
                "partially_worked": int((algorithm_frame["verdict"] == "Partially worked").sum()),
                "needs_review": int((algorithm_frame["verdict"] == "Needs review").sum()),
                "could_not_run": int((algorithm_frame["verdict"] == "Could not run").sum()),
            }
        )

    summary_frame = pd.DataFrame(summary_rows)
    best_overall = summary_frame.sort_values(["avg_score", "worked_well"], ascending=False).iloc[0]["algorithm"] if not summary_frame.empty else "N/A"
    avg_score = round(float(detail_frame["score"].mean())) if not detail_frame.empty else 0
    return sentences, sentence_reports, detail_frame, summary_frame, best_overall, avg_score, len(detailed_rows)


def verdict_class(verdict: str) -> str:
    verdict_lower = verdict.lower()
    if "worked" in verdict_lower:
        return "good"
    if "partial" in verdict_lower:
        return "review"
    if "could not" in verdict_lower or "unavailable" in verdict_lower:
        return "error"
    return "review"


def build_excel_report(input_text: str, sentence_reports: list[dict], detail_frame: pd.DataFrame, summary_frame: pd.DataFrame) -> bytes:
    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        sheet_name = "Sentence Details"
        detail_frame.to_excel(writer, index=False, sheet_name=sheet_name, startrow=3)
        worksheet = writer.sheets[sheet_name]

        title_fill = PatternFill("solid", fgColor="07111F")
        header_fill = PatternFill("solid", fgColor="17324E")
        good_fill = PatternFill("solid", fgColor="66E3A8")
        review_fill = PatternFill("solid", fgColor="F6C667")
        error_fill = PatternFill("solid", fgColor="FF7A7A")
        winner_fill = PatternFill("solid", fgColor="2C3E50")

        worksheet.merge_cells("A1:I1")
        worksheet["A1"] = "Sentence Segmentation Comparison Report"
        worksheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
        worksheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        worksheet["A1"].fill = title_fill

        worksheet.merge_cells("A2:I2")
        worksheet["A2"] = input_text
        worksheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        worksheet["A2"].fill = PatternFill("solid", fgColor="0B1628")

        worksheet.row_dimensions[1].height = 24
        worksheet.row_dimensions[2].height = 42
        worksheet.row_dimensions[4].height = 24

        for cell in worksheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        verdict_column = None
        winner_column = None
        for index, cell in enumerate(worksheet[4], start=1):
            if cell.value == "verdict":
                verdict_column = index
            if cell.value == "is_winner":
                winner_column = index

        for row in worksheet.iter_rows(min_row=5, max_row=worksheet.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            if verdict_column is not None:
                verdict_cell = row[verdict_column - 1]
                verdict_value = str(verdict_cell.value or "").lower()
                if "worked" in verdict_value:
                    verdict_cell.fill = good_fill
                elif "partial" in verdict_value:
                    verdict_cell.fill = review_fill
                elif "could not" in verdict_value or "unavailable" in verdict_value:
                    verdict_cell.fill = error_fill
            if winner_column is not None and str(row[winner_column - 1].value).lower() == "true":
                for cell in row:
                    cell.fill = winner_fill

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                cell_value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(cell_value))
            worksheet.column_dimensions[column_letter].width = min(max_length + 4, 48)

        worksheet.freeze_panes = "A5"
        worksheet.auto_filter.ref = worksheet.dimensions

        summary_sheet = writer.book.create_sheet("Algorithm Summary")
        summary_sheet.merge_cells("A1:F1")
        summary_sheet["A1"] = "Algorithm Summary"
        summary_sheet["A1"].font = Font(bold=True, color="FFFFFF", size=14)
        summary_sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
        summary_sheet["A1"].fill = title_fill
        summary_sheet.merge_cells("A2:F2")
        summary_sheet["A2"] = "Average score and verdict counts across all uploaded sentences."
        summary_sheet["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        summary_sheet["A2"].fill = PatternFill("solid", fgColor="0B1628")
        summary_sheet.row_dimensions[1].height = 24
        summary_sheet.row_dimensions[2].height = 30
        summary_frame.to_excel(writer, index=False, sheet_name="Algorithm Summary", startrow=3)
        for cell in summary_sheet[4]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in summary_sheet.iter_rows(min_row=5, max_row=summary_sheet.max_row):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        for column_cells in summary_sheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                cell_value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, len(cell_value))
            summary_sheet.column_dimensions[column_letter].width = min(max_length + 4, 40)
        summary_sheet.freeze_panes = "A5"
        summary_sheet.auto_filter.ref = summary_sheet.dimensions

    buffer.seek(0)
    return buffer.getvalue()


def build_table(results: list[dict], winner_algorithm: str) -> str:
    rows = []
    for item in results:
        badge_class = verdict_class(item["verdict"])
        is_winner = item["algorithm"] == winner_algorithm
        row_class = "winner-row" if is_winner else ""
        winner_badge = '<span class="winner-chip">Winner</span>' if is_winner else ''
        result_text = " <span class=\"separator\">|</span> ".join(
            html.escape(part) for part in item["result"].split(" | ")
        )
        reason_text = html.escape(item["reason"]).replace("\n", "<br>")
        rows.append(
            f"""
            <tr class=\"{row_class}\">
                <td>{html.escape(item["algorithm"])} {winner_badge}</td>
                <td class=\"result-cell\">{result_text}</td>
                <td>
                    <div class=\"badge {badge_class}\">{html.escape(item["verdict"])} </div>
                    <div class=\"reason-text\">{reason_text}</div>
                </td>
            </tr>
            """
        )

    return f"""
    <style>
        html, body {{
            margin: 0;
            padding: 0;
            background: transparent;
            font-family: Inter, ui-sans-serif, system-ui, -apple-system, BlinkMacSystemFont, \"Segoe UI\", sans-serif;
        }}

        .table-shell {{
            padding: 0.25rem 0.1rem 0.5rem;
        }}

        .result-table {{
            width: 100%;
            border-collapse: separate;
            border-spacing: 0 12px;
        }}

        .result-table th {{
            color: #a8b4cc;
            text-align: left;
            padding: 0 1rem 0.55rem 1rem;
            font-size: 0.8rem;
            letter-spacing: 0.08em;
            text-transform: uppercase;
        }}

        .result-table td {{
            padding: 0.95rem 1rem;
            vertical-align: top;
            color: #edf2ff;
            background: rgba(255,255,255,0.04);
            border-top: 1px solid rgba(255,255,255,0.08);
            border-bottom: 1px solid rgba(255,255,255,0.08);
        }}

        .result-table td:first-child {{
            border-left: 1px solid rgba(255,255,255,0.08);
            border-top-left-radius: 18px;
            border-bottom-left-radius: 18px;
            font-weight: 700;
            width: 16%;
        }}

        .result-table td:last-child {{
            border-right: 1px solid rgba(255,255,255,0.08);
            border-top-right-radius: 18px;
            border-bottom-right-radius: 18px;
            width: 34%;
        }}

        .result-cell {{
            color: #f5f8ff;
            line-height: 1.45;
            word-break: break-word;
            font-size: 0.95rem;
        }}

        .separator {{
            color: rgba(246, 198, 103, 0.95);
            font-weight: 700;
        }}

        .winner-chip {{
            display: inline-flex;
            margin-left: 0.45rem;
            padding: 0.18rem 0.55rem;
            border-radius: 999px;
            background: rgba(246, 198, 103, 0.18);
            border: 1px solid rgba(246, 198, 103, 0.38);
            color: #f6c667;
            font-size: 0.7rem;
            font-weight: 800;
            text-transform: uppercase;
            letter-spacing: 0.08em;
            vertical-align: middle;
        }}

        .winner-row td {{
            background: rgba(246, 198, 103, 0.07);
            border-top-color: rgba(246, 198, 103, 0.24);
            border-bottom-color: rgba(246, 198, 103, 0.24);
        }}

        .winner-row td:first-child {{
            box-shadow: inset 0 0 0 1px rgba(246, 198, 103, 0.18);
        }}

        .badge {{
            display: inline-flex;
            align-items: center;
            gap: 0.4rem;
            padding: 0.32rem 0.72rem;
            border-radius: 999px;
            font-size: 0.8rem;
            font-weight: 700;
            white-space: nowrap;
        }}

        .badge.good {{
            color: #052415;
            background: rgba(102, 227, 168, 0.95);
        }}

        .badge.review {{
            color: #2d1d05;
            background: rgba(246, 198, 103, 0.95);
        }}

        .badge.error {{
            color: #2d0909;
            background: rgba(255, 122, 122, 0.95);
        }}

        .reason-text {{
            margin-top: 0.7rem;
            color: #d7def0;
            line-height: 1.5;
            font-size: 0.92rem;
        }}

        .result-table tbody tr:hover td {{
            background: rgba(255,255,255,0.06);
        }}
    </style>
    <div class="table-shell">
        <table class="result-table">
            <thead>
                <tr>
                    <th>Algorithm</th>
                    <th>Result</th>
                    <th>Verdict and Reason</th>
                </tr>
            </thead>
            <tbody>
                {''.join(rows)}
            </tbody>
        </table>
    </div>
    """

def main() -> None:
    inject_styles()

    st.markdown(
        """
        <div class="hero">
            <span class="eyebrow">Sentence Segmentation Studio</span>
            <h1>Upload a test file and compare every algorithm in one view.</h1>
            <p>
                See how Rule-Based, Regex, UrduHack, Stanza, Dataset-Trained Punkt, Groq LLM, and Hybrid segmentation behave on the same text,
                then review the output, verdict, and explanation in a single polished dashboard.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.write("")

    with st.sidebar:
        st.markdown("### Controls")
        st.caption("Use the text area or upload a .txt file.")
        input_text = load_sentence()
        analyze_clicked = st.button("Analyze Sentence", use_container_width=True)
        st.markdown("---")
        st.markdown("### What this app does")
        st.write("Runs all segmentation algorithms on the uploaded test file, including a model trained on Urdu tweets.")
        st.write("Highlights whether each one matched the reference split.")
        st.write("Explains over-segmentation, under-segmentation, and missing models.")
        st.write("Can also use a dataset-trained Punkt model and a Groq LLM if the required dependencies or key are configured.")

    if not analyze_clicked and input_text:
        analyze_clicked = True

    if not input_text:
        st.info("Paste or upload a sentence to see the comparison table.")
        return

    if not analyze_clicked:
        st.stop()

    sentences, sentence_reports, detail_frame, summary_frame, best_overall, avg_score, total_runs = analyze_input_text(input_text)
    total_sentences = len(sentences)
    working = int((detail_frame["verdict"] == "Worked well").sum()) if not detail_frame.empty else 0
    review = int((detail_frame["verdict"] == "Needs review").sum()) if not detail_frame.empty else 0
    unavailable = int((detail_frame["verdict"] == "Could not run").sum()) if not detail_frame.empty else 0
    best_item = summary_frame.sort_values(["avg_score", "worked_well"], ascending=False).iloc[0] if not summary_frame.empty else None
    best_name = best_item["algorithm"] if best_item is not None else "N/A"
    best_score = int(round(best_item["avg_score"])) if best_item is not None else 0
    excel_bytes = build_excel_report(input_text, sentence_reports, detail_frame, summary_frame)

    col1, col2, col3, col4 = st.columns(4)
    col1.markdown(
        f"""
        <div class="metric-card">
            <div class="metric-label">Sentences</div>
            <div class="metric-value">{total_sentences}</div>
            <div class="metric-subtext">Processed in one run</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    col2.markdown(
        f"""
        <div class="metric-card">
            <div class="metric-label">Worked Well</div>
            <div class="metric-value">{working}</div>
            <div class="metric-subtext">Across all sentence runs</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    col3.markdown(
        f"""
        <div class="metric-card">
            <div class="metric-label">Needs Review</div>
            <div class="metric-value">{review}</div>
            <div class="metric-subtext">Across all sentence runs</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    col4.markdown(
        f"""
        <div class="metric-card">
            <div class="metric-label">Best Overall</div>
            <div class="winner-pill">Winner</div>
            <div class="metric-value" style="font-size:1.35rem;">{html.escape(best_name)}</div>
            <div class="metric-subtext">Score: {best_score}% | Average: {avg_score}% | {unavailable} unavailable</div>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.download_button(
        "Download Excel results",
        data=excel_bytes,
        file_name="sentence_segmentation_results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )

    st.write("")

    st.markdown(
        """
        <div class="panel">
            <h3 class="section-title">Sentence-by-sentence analysis</h3>
            <p style="margin:0;color:#a8b4cc;">Each sentence is analyzed separately, and every algorithm is applied in the same run.</p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    for report in sentence_reports:
        title = f"Sentence {report['sentence_index']}"
        subtitle = report['sentence'][:90] + ("..." if len(report['sentence']) > 90 else "")
        with st.expander(f"{title}: {subtitle}", expanded=report['sentence_index'] == 1):
            table_left, table_right = st.columns([1.5, 1])
            with table_left:
                components.html(build_table(report["results"], report["best_algorithm"]), height=220 + (len(report["results"]) * 155), scrolling=True)
            with table_right:
                st.markdown(
                    f"""
                    <div class="panel">
                        <h3 class="section-title">Sentence {report['sentence_index']}</h3>
                    """,
                    unsafe_allow_html=True,
                )
                st.code(report["sentence"], language="text")
                st.markdown(
                    f"""
                    <div class="footer-note">
                        Reference split count: <strong>{len(reference_split(report['sentence']))}</strong><br>
                        Winner for this sentence: <strong>{html.escape(report['best_algorithm'])}</strong><br>
                        Worked well: <strong>{report['worked_well']}</strong> | Needs review: <strong>{report['needs_review']}</strong> | Unavailable: <strong>{report['unavailable']}</strong>
                    </div>
                    """,
                    unsafe_allow_html=True,
                )
                st.markdown("</div>", unsafe_allow_html=True)

    st.caption(
        "Tip: the app now processes every sentence in the uploaded file separately and exports the full sentence-level report to Excel."
    )

    st.caption(
        "Tip: Stanza downloads its Urdu tokenizer the first time it runs. If the download is blocked, the row will explain the failure."
    )

    st.caption(
        "Dataset-Trained Punkt is learned from 50 cleaned Urdu tweets at startup. Groq LLM runs only when GROQ_API_KEY is set in your environment. You can optionally set GROQ_MODEL, with llama-3.1-8b-instant as the default."
    )


if __name__ == "__main__":
    main()











